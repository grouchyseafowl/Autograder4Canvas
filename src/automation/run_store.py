"""
RunStore — SQLite persistence layer for Academic Integrity Check results.

Replaces FlagAggregator (Excel-based, lossy, drops moderate/low/none concern levels).
Stores ALL concern levels with raw marker counts for retroactive re-weighting at
display time — no need to re-hit Canvas API to apply a different population profile.

DB lives at: <output_dir>/aic_runs.db

Invocation paths:
  1. automation_engine._run_adc()  — AIC runs as part of full grading workflow
  2. analyze_assignment() directly  — standalone AIC invoked from GUI or CLI
  save_result() must be called from both.
"""

import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, TYPE_CHECKING

if TYPE_CHECKING:
    from Academic_Dishonesty_Check_v2 import AnalysisResult


# ──────────────────────────────────────────────────────────────────────────────
# DB path helper
# ──────────────────────────────────────────────────────────────────────────────

def get_default_db_path() -> Path:
    """Return the path to aic_runs.db in the configured output directory."""
    try:
        import sys
        src_dir = Path(__file__).parent.parent
        if str(src_dir) not in sys.path:
            sys.path.insert(0, str(src_dir))
        from autograder_utils import get_output_base_dir
        base = get_output_base_dir()
    except Exception:
        base = Path.home() / "Canvas Grading"
    base.mkdir(parents=True, exist_ok=True)
    return base / "aic_runs.db"


# ──────────────────────────────────────────────────────────────────────────────
# DDL — all tables created idempotently
# ──────────────────────────────────────────────────────────────────────────────

_DDL = """
PRAGMA journal_mode = WAL;
PRAGMA foreign_keys = ON;

-- ─────────────────────────────────────────────────────────────────────────────
-- aic_results: one row per (student, assignment), upserted on re-run
-- ─────────────────────────────────────────────────────────────────────────────
CREATE TABLE IF NOT EXISTS aic_results (
    -- Identity key
    student_id          TEXT NOT NULL,
    assignment_id       TEXT NOT NULL,

    -- Context
    student_name        TEXT NOT NULL,
    course_id           TEXT NOT NULL,
    course_name         TEXT NOT NULL,
    assignment_name     TEXT NOT NULL,
    context_profile     TEXT NOT NULL DEFAULT 'standard',

    -- Timing
    submitted_at        TEXT,           -- ISO-8601; nullable (not always available)
    last_analyzed_at    TEXT NOT NULL,  -- set on every upsert

    -- Text stats
    word_count          INTEGER NOT NULL DEFAULT 0,

    -- Scores
    suspicious_score            REAL NOT NULL DEFAULT 0.0,
    authenticity_score          REAL NOT NULL DEFAULT 0.0,
    adjusted_suspicious_score   REAL,
    ai_organizational_score     REAL NOT NULL DEFAULT 0.0,

    -- Concern
    concern_level           TEXT NOT NULL DEFAULT 'none',
    adjusted_concern_level  TEXT,

    -- Human presence
    human_presence_confidence   REAL,
    human_presence_level        TEXT,
    human_presence_details      TEXT,   -- JSON dict

    -- Smoking gun (chatbot paste artifacts)
    smoking_gun         INTEGER NOT NULL DEFAULT 0,    -- 0 or 1
    smoking_gun_details TEXT    NOT NULL DEFAULT '[]', -- JSON array of detail strings

    -- Raw marker counts (JSON dict) — stored for retroactive profile re-weighting.
    -- Recalculation is local Python math (counts × profile multipliers); never
    -- requires re-hitting Canvas API.
    marker_counts       TEXT NOT NULL DEFAULT '{}',

    -- Outlier flags
    is_outlier          INTEGER NOT NULL DEFAULT 0,    -- 0 or 1
    outlier_reasons     TEXT    NOT NULL DEFAULT '[]', -- JSON array

    -- Guidance / conversation context (JSON arrays)
    context_adjustments     TEXT NOT NULL DEFAULT '[]',
    conversation_starters   TEXT NOT NULL DEFAULT '[]',
    verification_questions  TEXT NOT NULL DEFAULT '[]',
    revision_guidance       TEXT NOT NULL DEFAULT '[]',

    PRIMARY KEY (student_id, assignment_id)
);

CREATE INDEX IF NOT EXISTS idx_results_course
    ON aic_results (course_id, assignment_id);
CREATE INDEX IF NOT EXISTS idx_results_student_course
    ON aic_results (student_id, course_id);

-- ─────────────────────────────────────────────────────────────────────────────
-- teacher_notes: freeform conversation log
-- Does NOT affect scores — entirely separate from profile overrides.
-- ─────────────────────────────────────────────────────────────────────────────
CREATE TABLE IF NOT EXISTS teacher_notes (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    student_id      TEXT NOT NULL,
    course_id       TEXT NOT NULL,
    assignment_id   TEXT,          -- nullable: note may be global to the student
    note_text       TEXT NOT NULL,
    created_at      TEXT NOT NULL, -- ISO-8601
    updated_at      TEXT NOT NULL  -- ISO-8601
);

CREATE INDEX IF NOT EXISTS idx_notes_student
    ON teacher_notes (student_id, course_id);

-- ─────────────────────────────────────────────────────────────────────────────
-- student_profile_overrides: per-student population profile
-- Applies standardised multipliers to displayed scores.
-- UI must show "Scores recalculated with [ESL profile]" prominently.
-- ─────────────────────────────────────────────────────────────────────────────
CREATE TABLE IF NOT EXISTS student_profile_overrides (
    student_id  TEXT PRIMARY KEY,
    profile_id  TEXT NOT NULL,  -- 'standard'|'esl'|'neurodivergent'|'first_gen'|'community_college'
    set_at      TEXT NOT NULL   -- ISO-8601
);

-- ─────────────────────────────────────────────────────────────────────────────
-- grading_runs: one row per (course, assignment) grading pass
-- ─────────────────────────────────────────────────────────────────────────────
CREATE TABLE IF NOT EXISTS grading_runs (
  course_id           TEXT NOT NULL,
  assignment_id       TEXT NOT NULL,
  course_name         TEXT,
  assignment_name     TEXT,
  graded_at           TEXT NOT NULL,
  grading_tool        TEXT NOT NULL,
  total_students      INTEGER DEFAULT 0,
  complete_count      INTEGER DEFAULT 0,
  incomplete_count    INTEGER DEFAULT 0,
  skipped_count       INTEGER DEFAULT 0,
  flagged_count       INTEGER DEFAULT 0,
  min_word_count      INTEGER,
  was_dry_run         INTEGER DEFAULT 0,
  mode_settings       TEXT DEFAULT '{}',
  PRIMARY KEY (course_id, assignment_id)
);

-- ─────────────────────────────────────────────────────────────────────────────
-- grading_results: one row per (student, assignment) grading result
-- ─────────────────────────────────────────────────────────────────────────────
CREATE TABLE IF NOT EXISTS grading_results (
  student_id          TEXT NOT NULL,
  assignment_id       TEXT NOT NULL,
  student_name        TEXT NOT NULL,
  course_id           TEXT,
  course_name         TEXT,
  assignment_name     TEXT,
  submitted_at        TEXT,
  graded_at           TEXT NOT NULL,
  grade               TEXT NOT NULL,
  reason              TEXT NOT NULL,
  word_count          INTEGER DEFAULT 0,
  submission_type     TEXT,
  submission_body     TEXT,
  attachment_meta     TEXT DEFAULT '[]',
  flags               TEXT DEFAULT '[]',
  is_flagged          INTEGER DEFAULT 0,
  grading_tool        TEXT DEFAULT 'ci',
  min_word_count      INTEGER,
  was_skipped         INTEGER DEFAULT 0,
  post_count          INTEGER,
  reply_count         INTEGER,
  avg_words_per_post  REAL,
  teacher_override    TEXT,
  override_reason     TEXT,
  override_at         TEXT,
  PRIMARY KEY (student_id, assignment_id)
);

CREATE INDEX IF NOT EXISTS idx_grading_results_course
    ON grading_results (course_id, assignment_id);
"""


# ──────────────────────────────────────────────────────────────────────────────
# RunStore
# ──────────────────────────────────────────────────────────────────────────────

class RunStore:
    """
    SQLite persistence layer for AIC results.

    Thread-safety: single connection with check_same_thread=False; WAL mode
    handles concurrent readers.  All writes are short transactions.
    """

    def __init__(self, db_path: Optional[Path] = None):
        self.db_path = db_path or get_default_db_path()
        self._conn = sqlite3.connect(
            str(self.db_path),
            check_same_thread=False,
            detect_types=sqlite3.PARSE_DECLTYPES,
        )
        self._conn.row_factory = sqlite3.Row
        self._migrate()

    def close(self) -> None:
        self._conn.close()

    # ── Private helpers ────────────────────────────────────────────────────

    def _migrate(self) -> None:
        """Create / update schema (idempotent)."""
        self._conn.executescript(_DDL)
        self._apply_migrations()
        self._conn.commit()

    def _apply_migrations(self) -> None:
        """
        Apply incremental schema migrations (idempotent).
        SQLite doesn't support IF NOT EXISTS on ALTER TABLE, so we try each
        migration and swallow the 'duplicate column' error.
        """
        # Phase 8: Two-axis weight system — store provenance with each result
        migrations = [
            # aic_results: record which education level + population produced the scores
            "ALTER TABLE aic_results ADD COLUMN education_level TEXT",
            "ALTER TABLE aic_results ADD COLUMN population_settings TEXT",  # JSON dict
            # student_profile_overrides: composable per-student overrides
            "ALTER TABLE student_profile_overrides ADD COLUMN esl_override TEXT",
            "ALTER TABLE student_profile_overrides ADD COLUMN first_gen_override TEXT",
            "ALTER TABLE student_profile_overrides ADD COLUMN neurodivergent_override INTEGER",
        ]
        for sql in migrations:
            try:
                self._conn.execute(sql)
            except Exception:
                pass  # Column already exists — ignore

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).isoformat()

    @staticmethod
    def _j(value: Any) -> str:
        """Serialize to JSON array string; None → '[]'."""
        if value is None:
            return "[]"
        return json.dumps(value, default=str)

    @staticmethod
    def _jd(value: Any) -> str:
        """Serialize to JSON dict string; None → '{}'."""
        if value is None:
            return "{}"
        return json.dumps(value, default=str)

    # ── Write ──────────────────────────────────────────────────────────────

    def save_result(
        self,
        result: "AnalysisResult",
        *,
        course_id: str,
        course_name: str,
        assignment_id: str,
        assignment_name: str,
        submitted_at: Optional[str] = None,
        context_profile: str = "standard",
    ) -> None:
        """
        Upsert one AnalysisResult.

        On conflict (student_id, assignment_id) the row is fully replaced so
        re-runs always reflect the most recent analysis.
        """
        now = self._now()
        self._conn.execute(
            """
            INSERT INTO aic_results (
                student_id, assignment_id,
                student_name, course_id, course_name, assignment_name, context_profile,
                submitted_at, last_analyzed_at,
                word_count,
                suspicious_score, authenticity_score, adjusted_suspicious_score,
                ai_organizational_score,
                concern_level, adjusted_concern_level,
                human_presence_confidence, human_presence_level, human_presence_details,
                smoking_gun, smoking_gun_details,
                marker_counts,
                is_outlier, outlier_reasons,
                context_adjustments, conversation_starters,
                verification_questions, revision_guidance,
                education_level, population_settings
            ) VALUES (
                ?,?,  ?,?,?,?,?,  ?,?,  ?,  ?,?,?,  ?,  ?,?,
                ?,?,?,  ?,?,  ?,  ?,?,  ?,?,?,?,  ?,?
            )
            ON CONFLICT (student_id, assignment_id) DO UPDATE SET
                student_name              = excluded.student_name,
                course_id                 = excluded.course_id,
                course_name               = excluded.course_name,
                assignment_name           = excluded.assignment_name,
                context_profile           = excluded.context_profile,
                submitted_at              = excluded.submitted_at,
                last_analyzed_at          = excluded.last_analyzed_at,
                word_count                = excluded.word_count,
                suspicious_score          = excluded.suspicious_score,
                authenticity_score        = excluded.authenticity_score,
                adjusted_suspicious_score = excluded.adjusted_suspicious_score,
                ai_organizational_score   = excluded.ai_organizational_score,
                concern_level             = excluded.concern_level,
                adjusted_concern_level    = excluded.adjusted_concern_level,
                human_presence_confidence = excluded.human_presence_confidence,
                human_presence_level      = excluded.human_presence_level,
                human_presence_details    = excluded.human_presence_details,
                smoking_gun               = excluded.smoking_gun,
                smoking_gun_details       = excluded.smoking_gun_details,
                marker_counts             = excluded.marker_counts,
                is_outlier                = excluded.is_outlier,
                outlier_reasons           = excluded.outlier_reasons,
                context_adjustments       = excluded.context_adjustments,
                conversation_starters     = excluded.conversation_starters,
                verification_questions    = excluded.verification_questions,
                revision_guidance         = excluded.revision_guidance,
                education_level           = excluded.education_level,
                population_settings       = excluded.population_settings
            """,
            (
                str(result.student_id), str(assignment_id),
                result.student_name, str(course_id), course_name,
                assignment_name, context_profile,
                submitted_at, now,
                result.word_count,
                result.suspicious_score,
                result.authenticity_score,
                result.adjusted_suspicious_score,
                result.ai_organizational_score,
                result.concern_level,
                result.adjusted_concern_level,
                result.human_presence_confidence,
                result.human_presence_level,
                self._jd(result.human_presence_details),
                int(result.smoking_gun),
                self._j(result.smoking_gun_details),
                self._jd(result.marker_counts),
                int(result.is_outlier),
                self._j(result.outlier_reasons),
                self._j(result.context_adjustments_applied),
                self._j(result.conversation_starters),
                self._j(result.verification_questions),
                self._j(result.revision_guidance),
                getattr(result, 'education_level', None),
                self._jd(getattr(result, 'population_settings', None)),
            ),
        )
        self._conn.commit()

    # ── Should re-analyze? ─────────────────────────────────────────────────

    def should_reanalyze(
        self,
        student_id: str,
        assignment_id: str,
        submitted_at: Optional[str],
    ) -> bool:
        """
        Return True if the submission should be (re-)analyzed.

        True when:
        - No existing row for (student_id, assignment_id)
        - submitted_at is newer than the stored submitted_at (resubmission)
        """
        row = self._conn.execute(
            "SELECT submitted_at FROM aic_results "
            "WHERE student_id = ? AND assignment_id = ?",
            (str(student_id), str(assignment_id)),
        ).fetchone()

        if row is None:
            return True  # never analyzed

        if submitted_at is None:
            return False  # no timestamp to compare

        stored = row["submitted_at"]
        if stored is None:
            return True  # was analyzed without timestamp; re-run now we have one

        return submitted_at > stored  # newer submission → re-run

    # ── Read: run browser (sidebar) ────────────────────────────────────────

    def get_runs(self) -> List[Dict]:
        """
        Return distinct (course, assignment) run groups for the sidebar.

        Columns: course_id, course_name, assignment_id, assignment_name,
                 analyzed_count, smoking_gun_count, last_run
        """
        rows = self._conn.execute(
            """
            SELECT
                course_id,
                course_name,
                assignment_id,
                assignment_name,
                COUNT(*)              AS analyzed_count,
                SUM(smoking_gun)      AS smoking_gun_count,
                MAX(last_analyzed_at) AS last_run
            FROM aic_results
            GROUP BY course_id, assignment_id
            ORDER BY last_run DESC
            """
        ).fetchall()
        return [dict(r) for r in rows]

    # ── Read: cohort landscape (scatter plot) ──────────────────────────────

    def get_cohort(
        self,
        course_id: str,
        assignment_id: Optional[str] = None,
    ) -> List[Dict]:
        """
        Per-student data for the class landscape scatter plot.

        assignment_id given → that specific run.
        assignment_id omitted → latest result per student across the course.
        Returns marker_counts decoded from JSON.
        """
        if assignment_id:
            rows = self._conn.execute(
                """
                SELECT
                    student_id, student_name,
                    suspicious_score, authenticity_score,
                    adjusted_suspicious_score, adjusted_concern_level,
                    concern_level, ai_organizational_score,
                    human_presence_confidence, human_presence_level,
                    smoking_gun, word_count, last_analyzed_at,
                    marker_counts
                FROM aic_results
                WHERE course_id = ? AND assignment_id = ?
                ORDER BY suspicious_score DESC
                """,
                (str(course_id), str(assignment_id)),
            ).fetchall()
        else:
            rows = self._conn.execute(
                """
                SELECT
                    r.student_id, r.student_name,
                    r.suspicious_score, r.authenticity_score,
                    r.adjusted_suspicious_score, r.adjusted_concern_level,
                    r.concern_level, r.ai_organizational_score,
                    r.human_presence_confidence, r.human_presence_level,
                    r.smoking_gun, r.word_count, r.last_analyzed_at,
                    r.marker_counts
                FROM aic_results r
                INNER JOIN (
                    SELECT student_id, MAX(last_analyzed_at) AS latest
                    FROM aic_results WHERE course_id = ?
                    GROUP BY student_id
                ) latest_run
                  ON  r.student_id        = latest_run.student_id
                  AND r.last_analyzed_at  = latest_run.latest
                  AND r.course_id         = ?
                ORDER BY r.suspicious_score DESC
                """,
                (str(course_id), str(course_id)),
            ).fetchall()

        result = []
        for r in rows:
            d = dict(r)
            try:
                d["marker_counts"] = json.loads(d.get("marker_counts") or "{}")
            except (json.JSONDecodeError, TypeError):
                d["marker_counts"] = {}
            result.append(d)
        return result

    # ── Read: student trajectory (sparklines) ─────────────────────────────

    def get_trajectory(self, student_id: str, course_id: str) -> List[Dict]:
        """
        All analyzed submissions for one student in one course, ordered
        chronologically.  Used to render sparklines across assignments.

        Columns include: assignment_name, submitted_at, word_count,
        scores, concern_level, human_presence_confidence, smoking_gun.
        """
        rows = self._conn.execute(
            """
            SELECT
                assignment_id, assignment_name,
                submitted_at, last_analyzed_at,
                word_count,
                suspicious_score, authenticity_score,
                adjusted_suspicious_score, adjusted_concern_level,
                concern_level,
                human_presence_confidence, human_presence_level,
                smoking_gun, ai_organizational_score
            FROM aic_results
            WHERE student_id = ? AND course_id = ?
            ORDER BY COALESCE(submitted_at, last_analyzed_at) ASC
            """,
            (str(student_id), str(course_id)),
        ).fetchall()
        return [dict(r) for r in rows]

    # ── Read: student detail ───────────────────────────────────────────────

    def get_student_detail(
        self, student_id: str, assignment_id: str
    ) -> Optional[Dict]:
        """Full result row for one student + assignment, with JSON decoded."""
        row = self._conn.execute(
            "SELECT * FROM aic_results WHERE student_id = ? AND assignment_id = ?",
            (str(student_id), str(assignment_id)),
        ).fetchone()
        if row is None:
            return None
        d = dict(row)
        for col in (
            "human_presence_details",
            "smoking_gun_details",
            "marker_counts",
            "outlier_reasons",
            "context_adjustments",
            "conversation_starters",
            "verification_questions",
            "revision_guidance",
        ):
            raw = d.get(col)
            if raw:
                try:
                    d[col] = json.loads(raw)
                except (json.JSONDecodeError, TypeError):
                    pass
        return d

    def get_submission_content(
        self, student_id: str, assignment_id: str
    ) -> Optional[Dict]:
        """Fetch submission body and attachment metadata from grading_results."""
        row = self._conn.execute(
            "SELECT submission_body, submission_type, attachment_meta "
            "FROM grading_results WHERE student_id = ? AND assignment_id = ?",
            (str(student_id), str(assignment_id)),
        ).fetchone()
        if row is None:
            return None
        d = dict(row)
        raw = d.get("attachment_meta")
        if raw:
            try:
                d["attachment_meta"] = json.loads(raw)
            except (json.JSONDecodeError, TypeError):
                d["attachment_meta"] = []
        return d

    # ── Teacher notes ──────────────────────────────────────────────────────

    def save_note(
        self,
        student_id: str,
        course_id: str,
        note_text: str,
        assignment_id: Optional[str] = None,
    ) -> int:
        """Insert a new teacher note. Returns the new note id."""
        now = self._now()
        cur = self._conn.execute(
            """
            INSERT INTO teacher_notes
                (student_id, course_id, assignment_id, note_text, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (str(student_id), str(course_id), assignment_id, note_text, now, now),
        )
        self._conn.commit()
        return cur.lastrowid

    def update_note(self, note_id: int, note_text: str) -> None:
        """Update the text of an existing teacher note."""
        self._conn.execute(
            "UPDATE teacher_notes SET note_text = ?, updated_at = ? WHERE id = ?",
            (note_text, self._now(), note_id),
        )
        self._conn.commit()

    def delete_note(self, note_id: int) -> None:
        self._conn.execute("DELETE FROM teacher_notes WHERE id = ?", (note_id,))
        self._conn.commit()

    def get_notes(
        self,
        student_id: str,
        course_id: Optional[str] = None,
    ) -> List[Dict]:
        """Return teacher notes for a student, newest first."""
        if course_id:
            rows = self._conn.execute(
                "SELECT * FROM teacher_notes "
                "WHERE student_id = ? AND course_id = ? "
                "ORDER BY created_at DESC",
                (str(student_id), str(course_id)),
            ).fetchall()
        else:
            rows = self._conn.execute(
                "SELECT * FROM teacher_notes WHERE student_id = ? "
                "ORDER BY created_at DESC",
                (str(student_id),),
            ).fetchall()
        return [dict(r) for r in rows]

    # ── Profile overrides ──────────────────────────────────────────────────

    def set_profile_override(self, student_id: str, profile_id: str) -> None:
        """Set (or replace) a per-student population profile override."""
        self._conn.execute(
            """
            INSERT INTO student_profile_overrides (student_id, profile_id, set_at)
            VALUES (?, ?, ?)
            ON CONFLICT (student_id) DO UPDATE SET
                profile_id = excluded.profile_id,
                set_at     = excluded.set_at
            """,
            (str(student_id), profile_id, self._now()),
        )
        self._conn.commit()

    def set_composable_overrides(
        self,
        student_id: str,
        esl_level: str = "none",
        first_gen_level: str = "none",
        neurodivergent_aware: bool = False,
    ) -> None:
        """
        Save per-student composable population overrides (Phase 8 API).

        These are merged with class-level settings via max() when composing
        weights — student always gets the more protective setting.
        Also updates the legacy profile_id for backward compat.
        """
        # Derive a legacy profile_id string for backward compat display
        if neurodivergent_aware:
            legacy_pid = "neurodivergent"
        elif esl_level in ("moderate", "high"):
            legacy_pid = "esl"
        elif first_gen_level in ("moderate", "high"):
            legacy_pid = "first_gen"
        else:
            legacy_pid = "standard"

        self._conn.execute(
            """
            INSERT INTO student_profile_overrides
                (student_id, profile_id, set_at,
                 esl_override, first_gen_override, neurodivergent_override)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT (student_id) DO UPDATE SET
                profile_id             = excluded.profile_id,
                set_at                 = excluded.set_at,
                esl_override           = excluded.esl_override,
                first_gen_override     = excluded.first_gen_override,
                neurodivergent_override = excluded.neurodivergent_override
            """,
            (
                str(student_id), legacy_pid, self._now(),
                esl_level, first_gen_level, int(neurodivergent_aware),
            ),
        )
        self._conn.commit()

    def get_composable_overrides(self, student_id: str) -> Dict[str, Any]:
        """
        Return per-student composable overrides dict, or defaults if not set.

        Returns:
            Dict with keys: esl_level, first_gen_level, neurodivergent_aware
        """
        row = self._conn.execute(
            "SELECT esl_override, first_gen_override, neurodivergent_override "
            "FROM student_profile_overrides WHERE student_id = ?",
            (str(student_id),),
        ).fetchone()
        if row and (row["esl_override"] or row["first_gen_override"]
                    or row["neurodivergent_override"] is not None):
            return {
                "esl_level": row["esl_override"] or "none",
                "first_gen_level": row["first_gen_override"] or "none",
                "neurodivergent_aware": bool(row["neurodivergent_override"]),
            }
        return {"esl_level": "none", "first_gen_level": "none", "neurodivergent_aware": False}

    def get_profile_override(self, student_id: str) -> Optional[str]:
        """Return the profile_id override for a student, or None if unset."""
        row = self._conn.execute(
            "SELECT profile_id FROM student_profile_overrides WHERE student_id = ?",
            (str(student_id),),
        ).fetchone()
        return row["profile_id"] if row else None

    # ── Grading results ───────────────────────────────────────────────────

    def save_grading_result(self, result_dict: Dict) -> None:
        """Upsert one grading result row. result_dict keys map to column names."""
        d = dict(result_dict)
        # JSON-serialize list/dict columns
        d["flags"] = self._j(d.get("flags"))
        d["attachment_meta"] = self._j(d.get("attachment_meta"))
        # Default graded_at to now if not provided
        if not d.get("graded_at"):
            d["graded_at"] = self._now()

        cols = [
            "student_id", "assignment_id", "student_name",
            "course_id", "course_name", "assignment_name",
            "submitted_at", "graded_at", "grade", "reason",
            "word_count", "submission_type", "submission_body",
            "attachment_meta", "flags", "is_flagged",
            "grading_tool", "min_word_count", "was_skipped",
            "post_count", "reply_count", "avg_words_per_post",
            "teacher_override", "override_reason", "override_at",
        ]
        vals = [d.get(c) for c in cols]
        placeholders = ", ".join("?" * len(cols))
        col_names = ", ".join(cols)
        updates = ", ".join(f"{c} = excluded.{c}" for c in cols
                            if c not in ("student_id", "assignment_id"))

        self._conn.execute(
            f"INSERT INTO grading_results ({col_names}) "
            f"VALUES ({placeholders}) "
            f"ON CONFLICT (student_id, assignment_id) DO UPDATE SET {updates}",
            vals,
        )
        self._conn.commit()

    def save_grading_run(self, run_dict: Dict) -> None:
        """Upsert run-level metadata for one (course, assignment) grading pass."""
        d = dict(run_dict)
        d["mode_settings"] = self._jd(d.get("mode_settings"))
        if not d.get("graded_at"):
            d["graded_at"] = self._now()

        cols = [
            "course_id", "assignment_id", "course_name", "assignment_name",
            "graded_at", "grading_tool", "total_students",
            "complete_count", "incomplete_count", "skipped_count",
            "flagged_count", "min_word_count", "was_dry_run", "mode_settings",
        ]
        vals = [d.get(c) for c in cols]
        placeholders = ", ".join("?" * len(cols))
        col_names = ", ".join(cols)
        updates = ", ".join(f"{c} = excluded.{c}" for c in cols
                            if c not in ("course_id", "assignment_id"))

        self._conn.execute(
            f"INSERT INTO grading_runs ({col_names}) "
            f"VALUES ({placeholders}) "
            f"ON CONFLICT (course_id, assignment_id) DO UPDATE SET {updates}",
            vals,
        )
        self._conn.commit()

    def get_grading_assignments(self) -> List[Dict]:
        """
        Returns assignments grouped by course with counts for sidebar.

        LEFT JOINs grading_runs with aggregate counts from grading_results
        and aic_results to power both amber (grading) and rose (AIC)
        sidebar indicators.
        """
        rows = self._conn.execute(
            """
            SELECT
                gr.course_id,
                gr.course_name,
                gr.assignment_id,
                gr.assignment_name,
                gr.graded_at,
                gr.grading_tool,
                gr.total_students,
                gr.complete_count,
                gr.incomplete_count,
                COALESCE(res.flagged_count, 0) AS flagged_count,
                COALESCE(aic.elevated_count, 0) AS aic_elevated_count,
                COALESCE(aic.smoking_gun_count, 0) AS aic_smoking_gun_count
            FROM grading_runs gr
            LEFT JOIN (
                SELECT
                    course_id, assignment_id,
                    SUM(is_flagged) AS flagged_count
                FROM grading_results
                GROUP BY course_id, assignment_id
            ) res
              ON gr.course_id = res.course_id
             AND gr.assignment_id = res.assignment_id
            LEFT JOIN (
                SELECT
                    course_id, assignment_id,
                    SUM(CASE WHEN concern_level IN ('elevated','high') THEN 1 ELSE 0 END)
                        AS elevated_count,
                    SUM(smoking_gun) AS smoking_gun_count
                FROM aic_results
                GROUP BY course_id, assignment_id
            ) aic
              ON gr.course_id = aic.course_id
             AND gr.assignment_id = aic.assignment_id
            ORDER BY gr.graded_at DESC
            """
        ).fetchall()
        return [dict(r) for r in rows]

    def get_grading_cohort(self, course_id: str, assignment_id: str) -> List[Dict]:
        """All students for one assignment, with JSON decoded."""
        rows = self._conn.execute(
            "SELECT * FROM grading_results "
            "WHERE course_id = ? AND assignment_id = ? "
            "ORDER BY student_name ASC",
            (str(course_id), str(assignment_id)),
        ).fetchall()
        result = []
        for r in rows:
            d = dict(r)
            for col in ("flags", "attachment_meta"):
                raw = d.get(col)
                if raw:
                    try:
                        d[col] = json.loads(raw)
                    except (json.JSONDecodeError, TypeError):
                        pass
            result.append(d)
        return result

    def get_grading_with_aic(
        self, course_id: str, assignment_id: str
    ) -> List[Dict]:
        """LEFT JOIN grading_results to aic_results for AIC overlay."""
        rows = self._conn.execute(
            """
            SELECT
                g.*,
                a.concern_level      AS aic_concern_level,
                a.suspicious_score   AS aic_suspicious_score,
                a.human_presence_confidence AS aic_human_presence_confidence,
                a.smoking_gun        AS aic_smoking_gun
            FROM grading_results g
            LEFT JOIN aic_results a
              ON g.student_id = a.student_id
             AND g.assignment_id = a.assignment_id
            WHERE g.course_id = ? AND g.assignment_id = ?
            ORDER BY g.student_name ASC
            """,
            (str(course_id), str(assignment_id)),
        ).fetchall()
        result = []
        for r in rows:
            d = dict(r)
            for col in ("flags", "attachment_meta"):
                raw = d.get(col)
                if raw:
                    try:
                        d[col] = json.loads(raw)
                    except (json.JSONDecodeError, TypeError):
                        pass
            result.append(d)
        return result

    def set_teacher_override(
        self,
        student_id: str,
        assignment_id: str,
        grade: str,
        reason: str,
    ) -> None:
        """Persist a teacher override for a grading result."""
        self._conn.execute(
            "UPDATE grading_results "
            "SET teacher_override = ?, override_reason = ?, override_at = ? "
            "WHERE student_id = ? AND assignment_id = ?",
            (grade, reason, self._now(), str(student_id), str(assignment_id)),
        )
        self._conn.commit()

    def export_grading_xlsx(
        self,
        course_id: str,
        assignment_id: str,
        output_path: str,
    ) -> str:
        """Generate XLSX from stored grading data. Returns the output path."""
        import openpyxl

        rows = self.get_grading_with_aic(course_id, assignment_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Grading Results"

        if not rows:
            ws.append(["No results found"])
            wb.save(output_path)
            return output_path

        # Column headers — grading fields + AIC overlay fields
        headers = [
            "Student ID", "Student Name", "Course ID", "Course Name",
            "Assignment ID", "Assignment Name", "Submitted At", "Graded At",
            "Grade", "Reason", "Word Count", "Submission Type",
            "Flags", "Is Flagged", "Grading Tool", "Min Word Count",
            "Was Skipped", "Post Count", "Reply Count",
            "Avg Words Per Post",
            "Teacher Override", "Override Reason", "Override At",
            "AIC Concern Level", "AIC Suspicious Score",
            "AIC Human Presence Confidence", "AIC Smoking Gun",
        ]
        ws.append(headers)

        key_map = [
            "student_id", "student_name", "course_id", "course_name",
            "assignment_id", "assignment_name", "submitted_at", "graded_at",
            "grade", "reason", "word_count", "submission_type",
            "flags", "is_flagged", "grading_tool", "min_word_count",
            "was_skipped", "post_count", "reply_count",
            "avg_words_per_post",
            "teacher_override", "override_reason", "override_at",
            "aic_concern_level", "aic_suspicious_score",
            "aic_human_presence_confidence", "aic_smoking_gun",
        ]

        for row in rows:
            values = []
            for k in key_map:
                v = row.get(k)
                # Serialize lists/dicts back to string for spreadsheet
                if isinstance(v, (list, dict)):
                    v = json.dumps(v, default=str)
                values.append(v)
            ws.append(values)

        wb.save(output_path)
        return output_path

    # ── Cleanup helpers ────────────────────────────────────────────────────

    def count_for_cleanup(
        self,
        days: int,
        include_aic: bool = True,
        include_grading: bool = True,
        include_notes: bool = False,
        include_profiles: bool = False,
    ) -> Dict[str, int]:
        """
        Dry-run count for cleanup dialog preview.

        days applies to aic_results (last_analyzed_at), grading_results
        (graded_at), and teacher_notes (created_at).  Profile overrides
        are counted in full (no age filter — they represent a full reset).

        Returns dict with keys: 'aic', 'grading', 'notes', 'profiles'.
        """
        cutoff = self._cutoff_iso(days)
        counts: Dict[str, int] = {"aic": 0, "grading": 0, "notes": 0, "profiles": 0}

        if include_aic:
            row = self._conn.execute(
                "SELECT COUNT(*) FROM aic_results WHERE last_analyzed_at < ?",
                (cutoff,),
            ).fetchone()
            counts["aic"] = row[0] if row else 0

        if include_grading:
            row = self._conn.execute(
                "SELECT COUNT(*) FROM grading_results WHERE graded_at < ?",
                (cutoff,),
            ).fetchone()
            counts["grading"] = row[0] if row else 0

        if include_notes:
            row = self._conn.execute(
                "SELECT COUNT(*) FROM teacher_notes WHERE created_at < ?",
                (cutoff,),
            ).fetchone()
            counts["notes"] = row[0] if row else 0

        if include_profiles:
            row = self._conn.execute(
                "SELECT COUNT(*) FROM student_profile_overrides",
            ).fetchone()
            counts["profiles"] = row[0] if row else 0

        return counts

    def delete_for_cleanup(
        self,
        days: int,
        include_aic: bool = True,
        include_grading: bool = True,
        include_notes: bool = False,
        include_profiles: bool = False,
    ) -> Dict[str, int]:
        """
        Permanently delete records matching the cleanup criteria.

        For grading: also removes orphaned grading_runs after deleting results.
        For profiles: deletes all overrides (full reset, no age filter).

        Returns dict with keys: 'aic', 'grading', 'notes', 'profiles'.
        """
        cutoff = self._cutoff_iso(days)
        deleted: Dict[str, int] = {"aic": 0, "grading": 0, "notes": 0, "profiles": 0}

        if include_aic:
            cur = self._conn.execute(
                "DELETE FROM aic_results WHERE last_analyzed_at < ?",
                (cutoff,),
            )
            deleted["aic"] = cur.rowcount

        if include_grading:
            cur = self._conn.execute(
                "DELETE FROM grading_results WHERE graded_at < ?",
                (cutoff,),
            )
            deleted["grading"] = cur.rowcount
            # Remove orphaned run metadata
            self._conn.execute(
                """
                DELETE FROM grading_runs
                WHERE NOT EXISTS (
                    SELECT 1 FROM grading_results gr
                    WHERE gr.course_id     = grading_runs.course_id
                      AND gr.assignment_id = grading_runs.assignment_id
                )
                AND graded_at < ?
                """,
                (cutoff,),
            )

        if include_notes:
            cur = self._conn.execute(
                "DELETE FROM teacher_notes WHERE created_at < ?",
                (cutoff,),
            )
            deleted["notes"] = cur.rowcount

        if include_profiles:
            cur = self._conn.execute("DELETE FROM student_profile_overrides")
            deleted["profiles"] = cur.rowcount

        self._conn.commit()
        return deleted

    @staticmethod
    def _cutoff_iso(days: int) -> str:
        """Return ISO-8601 timestamp for *days* ago (UTC)."""
        from datetime import timedelta
        cutoff = datetime.now(timezone.utc) - timedelta(days=days)
        return cutoff.isoformat()
