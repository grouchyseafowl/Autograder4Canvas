"""
ExportReportsDialog — export AIC and grading data to PDF, Excel, or CSV.

Format:
  PDF   — QTextDocument → QPdfWriter (no extra dependencies)
  Excel — openpyxl via cohort data from RunStore
  CSV   — standard csv module

Scope:
  course_id + course_name  → course-level export (all students)
  + student_id / name      → per-student export (scoped to one student)
"""

from __future__ import annotations

import csv
import os
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QFileDialog, QFrame, QMessageBox,
)

from gui.styles import (
    px,
    SPACING_SM, SPACING_MD, SPACING_LG,
    PHOSPHOR_HOT, PHOSPHOR_MID, PHOSPHOR_DIM,
    BG_VOID,
    make_section_label,
    make_secondary_button, make_run_button,
)
from gui.widgets.phosphor_chip import PhosphorChip
from gui.widgets.switch_toggle import SwitchToggle


_VALUE_CSS = (f"color: {PHOSPHOR_MID}; font-size: {px(11)}px;"
              f" background: transparent; border: none;")
_DIM_CSS   = (f"color: {PHOSPHOR_DIM}; font-size: {px(10)}px;"
              f" background: transparent; border: none;")


class ExportReportsDialog(QDialog):
    """Export AIC and grading data to PDF, Excel, or CSV."""

    def __init__(
        self,
        store=None,
        course_id: str = "",
        course_name: str = "",
        student_id: str = "",
        student_name: str = "",
        parent=None,
    ):
        super().__init__(parent)
        self._store       = store
        self._course_id   = str(course_id)
        self._course_name = course_name or "Unknown Course"
        self._student_id  = str(student_id)
        self._student_name = student_name or ""

        self._format = "pdf"

        self.setWindowTitle("Export Reports")
        self.setMinimumWidth(520)
        self.setStyleSheet(f"QDialog {{ background: {BG_VOID}; }}")
        self._build_ui()

    # ── UI ────────────────────────────────────────────────────────────────────

    @staticmethod
    def _amber_sep() -> QFrame:
        """Amber gradient separator — bright at center, fading at edges."""
        sep = QFrame()
        sep.setFixedHeight(1)
        sep.setStyleSheet("""
            QFrame {
                background: qlineargradient(x1:0,y1:0,x2:1,y2:0,
                    stop:0.00 rgba(240,168,48,0),
                    stop:0.20 rgba(240,168,48,0.35),
                    stop:0.50 rgba(240,168,48,0.70),
                    stop:0.80 rgba(240,168,48,0.35),
                    stop:1.00 rgba(240,168,48,0));
                border: none;
            }
        """)
        return sep

    def _build_ui(self) -> None:
        root = QVBoxLayout(self)
        root.setContentsMargins(SPACING_LG, SPACING_MD, SPACING_LG, SPACING_MD)
        root.setSpacing(SPACING_MD)

        # ── Title block ───────────────────────────────────────────────────────
        title = QLabel("EXPORT REPORTS")
        title.setStyleSheet(
            f"color: {PHOSPHOR_HOT}; font-size: {px(16)}px; font-weight: bold;"
            f" letter-spacing: 2px; background: transparent; border: none;"
        )
        root.addWidget(title)

        sub = QLabel("Snapshot exports from the AIC RunStore — read-only, non-destructive.")
        sub.setStyleSheet(
            f"color: {PHOSPHOR_DIM}; font-size: {px(11)}px;"
            f" background: transparent; border: none;"
        )
        sub.setWordWrap(True)
        root.addWidget(sub)

        root.addWidget(self._amber_sep())

        # ── Scope (at top) ────────────────────────────────────────────────────
        root.addWidget(make_section_label("Scope"))
        root.addSpacing(4)

        def _scope_row(label_text: str, value_text: str) -> QHBoxLayout:
            row = QHBoxLayout()
            row.setSpacing(SPACING_SM)
            lbl = QLabel(label_text)
            lbl.setStyleSheet(_DIM_CSS)
            row.addWidget(lbl)
            val = QLabel(value_text)
            val.setStyleSheet(_VALUE_CSS)
            row.addWidget(val)
            row.addStretch()
            return row

        root.addLayout(_scope_row("Course:", self._course_name or "All available data"))
        if self._student_name:
            root.addLayout(_scope_row("Student:", self._student_name))

        root.addWidget(self._amber_sep())

        # ── Format ────────────────────────────────────────────────────────────
        root.addWidget(make_section_label("Format"))
        root.addSpacing(4)

        fmt_row = QHBoxLayout()
        fmt_row.setSpacing(SPACING_SM)
        self._fmt_chips: List[tuple[PhosphorChip, str]] = []
        for label, key in [("PDF", "pdf"), ("Excel  .xlsx", "xlsx"), ("CSV", "csv")]:
            chip = PhosphorChip(label, active=(key == "pdf"))
            chip.toggled.connect(lambda checked, k=key: self._on_format_chip(k, checked))
            fmt_row.addWidget(chip)
            self._fmt_chips.append((chip, key))
        fmt_row.addStretch()
        root.addLayout(fmt_row)

        # ── Include section (PDF only — toggled show/hide) ────────────────────
        self._include_sep = self._amber_sep()
        root.addWidget(self._include_sep)

        self._include_frame = QFrame()
        self._include_frame.setStyleSheet(
            "QFrame { background: transparent; border: none; }"
        )
        inc = QVBoxLayout(self._include_frame)
        inc.setContentsMargins(0, 0, 0, 0)
        inc.setSpacing(SPACING_SM)

        inc.addWidget(make_section_label("Include in PDF"))
        inc.addSpacing(4)

        self._toggle_summary     = SwitchToggle("Class Summary",       wrap_width=200)
        self._toggle_per_student = SwitchToggle("Per-Student Reports", wrap_width=200)
        self._toggle_trajectory  = SwitchToggle("Trajectory Data",     wrap_width=200)
        self._toggle_summary.setChecked(True)
        self._toggle_per_student.setChecked(True)

        _desc_css = (f"color: {PHOSPHOR_DIM}; font-size: {px(10)}px;"
                     f" background: transparent; border: none; margin-left: {px(43)}px;")
        for toggle, desc in [
            (self._toggle_summary,     "Distribution table, conversation opportunity counts, smoking-gun list."),
            (self._toggle_per_student, "One section per flagged student with marker breakdown."),
            (self._toggle_trajectory,  "Submission history across all assignments for each student."),
        ]:
            inc.addWidget(toggle)
            d = QLabel(desc)
            d.setStyleSheet(_desc_css)
            d.setWordWrap(True)
            inc.addWidget(d)

        root.addWidget(self._include_frame)

        self._include_post_sep = self._amber_sep()
        root.addWidget(self._include_post_sep)

        root.addStretch()
        root.addWidget(self._amber_sep())

        # ── Footer — flat row, no QFrame wrapper ───────────────────────────────
        footer = QHBoxLayout()
        footer.setSpacing(SPACING_SM)
        footer.addStretch()
        cancel_btn = QPushButton("Cancel")
        make_secondary_button(cancel_btn)
        cancel_btn.setFixedWidth(90)
        cancel_btn.clicked.connect(self.reject)
        footer.addWidget(cancel_btn)
        self._export_btn = QPushButton("Export")
        make_run_button(self._export_btn)
        self._export_btn.setFixedWidth(110)
        self._export_btn.clicked.connect(self._on_export)
        footer.addWidget(self._export_btn)
        root.addLayout(footer)

    # ── Interactions ──────────────────────────────────────────────────────────

    def _on_format_chip(self, key: str, checked: bool) -> None:
        if not checked:
            if self._format == key:
                for chip, k in self._fmt_chips:
                    if k == key:
                        chip.blockSignals(True)
                        chip.setChecked(True)
                        chip.blockSignals(False)
            return
        self._format = key
        for chip, k in self._fmt_chips:
            if k != key:
                chip.blockSignals(True)
                chip.setChecked(False)
                chip.blockSignals(False)
        is_pdf = (key == "pdf")
        self._include_sep.setVisible(is_pdf)
        self._include_frame.setVisible(is_pdf)
        self._include_post_sep.setVisible(is_pdf)
        self.adjustSize()

    def _suggested_path(self) -> str:
        """Build a suggested save path based on format + course + timestamp."""
        ts   = datetime.now().strftime("%Y%m%d_%H%M")
        stem = self._safe_name(self._course_name or "export")
        if self._student_id:
            stem += f"_{self._safe_name(self._student_name or self._student_id)}"
        ext_map = {"pdf": ".pdf", "xlsx": ".xlsx", "csv": ".csv"}
        ext = ext_map.get(self._format, ".pdf")
        prefix = "AIC_Report_" if self._format == "pdf" else "AIC_"
        name = f"{prefix}{stem}_{ts}{ext}"
        return str(Path.home() / "Desktop" / name)

    def _on_export(self) -> None:
        if not self._store:
            QMessageBox.warning(self, "Export", "No data store available.")
            return

        # Native save dialog
        filter_map = {
            "pdf":  "PDF Files (*.pdf)",
            "xlsx": "Excel Files (*.xlsx)",
            "csv":  "CSV Files (*.csv)",
        }
        path, _ = QFileDialog.getSaveFileName(
            self,
            "Export Report",
            self._suggested_path(),
            filter_map.get(self._format, "All Files (*)"),
        )
        if not path:
            return

        try:
            if self._format == "pdf":
                self._export_pdf(path)
            elif self._format == "xlsx":
                self._export_xlsx(path)
            else:
                self._export_csv(path)

            QMessageBox.information(
                self, "Export Complete",
                f"Saved to:\n{os.path.basename(path)}\n\nin {os.path.dirname(path)}",
            )
            self.accept()
        except Exception as exc:
            QMessageBox.critical(self, "Export Failed", str(exc))

    # ── Export: Excel ─────────────────────────────────────────────────────────

    def _export_xlsx(self, out: str) -> None:
        import openpyxl

        wb = openpyxl.Workbook()

        # Sheet 1: AIC cohort
        ws = wb.active
        ws.title = "AIC Results"
        rows = self._get_cohort_rows()
        aic_headers = [
            "Student ID", "Student Name", "Engagement Depth", "Authenticity Score",
            "Conversation Opportunity", "Personal Connection", "Smoking Gun",
            "Word Count", "Last Analyzed",
        ]
        aic_keys = [
            "student_id", "student_name", "suspicious_score", "authenticity_score",
            "concern_level", "human_presence_confidence", "smoking_gun",
            "word_count", "last_analyzed_at",
        ]
        ws.append(aic_headers)
        for r in rows:
            ws.append([r.get(k) for k in aic_keys])

        # Sheet 2: Trajectory
        ws2 = wb.create_sheet("Trajectory")
        ws2.append([
            "Student ID", "Student Name", "Assignment", "Submitted At",
            "Word Count", "Engagement Depth", "Conversation Opportunity",
            "Personal Connection", "Smoking Gun",
        ])
        for sid, sname in self._student_ids_and_names(rows):
            for t in self._store.get_trajectory(sid, self._course_id):
                ws2.append([
                    sid, sname,
                    t.get("assignment_name"), t.get("submitted_at"),
                    t.get("word_count"), t.get("suspicious_score"),
                    t.get("concern_level"), t.get("human_presence_confidence"),
                    t.get("smoking_gun"),
                ])

        wb.save(out)

    # ── Export: CSV ───────────────────────────────────────────────────────────

    def _export_csv(self, out: str) -> None:
        rows = self._get_cohort_rows()
        headers = [
            "student_id", "student_name", "suspicious_score", "authenticity_score",
            "concern_level", "human_presence_confidence", "smoking_gun",
            "word_count", "last_analyzed_at",
        ]
        with open(out, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)

    # ── Export: PDF ───────────────────────────────────────────────────────────

    def _export_pdf(self, out: str) -> None:
        from PySide6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout
        from PySide6.QtCore import QMarginsF

        sections: List[str] = []
        is_first = True

        def add_section(html: str) -> None:
            nonlocal is_first
            if is_first:
                sections.append(html)
                is_first = False
            else:
                sections.append(
                    f'<div style="page-break-before: always;">{html}</div>'
                )

        if self._toggle_summary.isChecked():
            add_section(self._pdf_class_summary())
        if self._toggle_per_student.isChecked():
            add_section(self._pdf_per_student())
        if self._toggle_trajectory.isChecked():
            add_section(self._pdf_trajectory())

        if not sections:
            raise ValueError("Select at least one section to include.")

        doc = QTextDocument()
        doc.setHtml(_PDF_TEMPLATE.format(body="\n".join(sections)))

        writer = QPdfWriter(out)
        writer.setPageLayout(QPageLayout(
            QPageSize(QPageSize.PageSizeId.Letter),
            QPageLayout.Orientation.Portrait,
            QMarginsF(15, 15, 15, 15),
            QPageLayout.Unit.Millimeter,
        ))
        doc.print_(writer)

    # ── PDF section builders ──────────────────────────────────────────────────

    def _pdf_class_summary(self) -> str:
        rows = self._get_cohort_rows()
        total = len(rows)
        counts: Dict[str, int] = {}
        for r in rows:
            lvl = str(r.get("concern_level") or "none").lower()
            counts[lvl] = counts.get(lvl, 0) + 1
        sg_count = sum(1 for r in rows if r.get("smoking_gun"))
        pct = lambda n: f"{n / total * 100:.0f}%" if total else "—"

        dist_rows = "".join(
            f"<tr><td>{lvl.title()}</td><td>{counts.get(lvl, 0)}</td>"
            f"<td>{pct(counts.get(lvl, 0))}</td></tr>"
            for lvl in ("high", "moderate", "low", "none")
        )
        dist_rows += (f'<tr class="total"><td><b>Total</b></td>'
                      f'<td><b>{total}</b></td><td></td></tr>')

        flagged_rows = "".join(
            f"<tr><td>{r.get('student_name', '—')}</td>"
            f"<td>{r.get('concern_level', '—')}</td>"
            f"<td>{r.get('suspicious_score', '—')}</td>"
            f"<td>{'YES' if r.get('smoking_gun') else ''}</td></tr>"
            for r in sorted(rows, key=lambda x: x.get("suspicious_score", 0) or 0,
                            reverse=True)
            if (r.get("concern_level") or "none").lower() in ("high", "moderate")
        )

        date_str = datetime.now().strftime("%B %d, %Y")
        flagged_section = (
            f"<h2>Conversation Opportunities</h2>"
            f"<table><tr><th>Student</th><th>Opportunity</th>"
            f"<th>Engagement</th><th>Smoking Gun</th></tr>{flagged_rows}</table>"
            if flagged_rows else ""
        )
        return (
            f"<h1>Class Summary</h1>"
            f"<p class='meta'>Course: <b>{self._course_name}</b>"
            f" &nbsp;|&nbsp; Generated: {date_str}</p>"
            f"<h2>Distribution</h2>"
            f"<table><tr><th>Conversation Opportunity</th><th>Count</th><th>%</th></tr>"
            f"{dist_rows}</table>"
            f"<p>Smoking-gun submissions: <b>{sg_count}</b></p>"
            f"{flagged_section}"
        )

    def _pdf_per_student(self) -> str:
        rows = self._get_cohort_rows()
        if not self._student_id:
            rows = [r for r in rows
                    if (r.get("concern_level") or "none").lower()
                    in ("high", "moderate", "low")]

        sections = []
        for i, r in enumerate(rows):
            mc = r.get("marker_counts") or {}
            marker_rows = "".join(
                f"<tr><td>{k.replace('_', ' ').title()}</td><td>{v}</td></tr>"
                for k, v in mc.items() if v
            ) if isinstance(mc, dict) else ""
            marker_table = (
                f"<h2>Markers</h2><table>"
                f"<tr><th>Marker</th><th>Count</th></tr>"
                f"{marker_rows}</table>"
                if marker_rows else ""
            )
            sep = '<div style="page-break-before: always;">' if i > 0 else "<div>"
            sections.append(
                f"{sep}"
                f"<h1>{r.get('student_name', 'Unknown Student')}</h1>"
                f"<p class='meta'>"
                f"Conversation: <b>{r.get('concern_level', '—')}</b>"
                f" &nbsp;|&nbsp; Engagement Depth: {r.get('suspicious_score', '—')}"
                f" &nbsp;|&nbsp; Personal Connection: {r.get('human_presence_confidence', '—')}"
                + (" &nbsp;|&nbsp; <span class='flag'>⚑ Smoking Gun</span>"
                   if r.get("smoking_gun") else "")
                + f"</p>{marker_table}"
                f"<p class='note'>For teacher reference only — not for students.</p>"
                f"</div>"
            )

        return "\n".join(sections) if sections else (
            "<h1>Per-Student Reports</h1><p>No flagged students found.</p>"
        )

    def _pdf_trajectory(self) -> str:
        rows = self._get_cohort_rows()
        student_ids = (
            [(self._student_id, self._student_name)] if self._student_id
            else self._student_ids_and_names(rows)
        )

        sections = []
        for i, (sid, sname) in enumerate(student_ids):
            traj = self._store.get_trajectory(sid, self._course_id)
            if not traj:
                continue
            traj_rows = "".join(
                f"<tr><td>{t.get('assignment_name', '—')}</td>"
                f"<td>{t.get('submitted_at', '—')}</td>"
                f"<td>{t.get('suspicious_score', '—')}</td>"
                f"<td>{t.get('concern_level', '—')}</td>"
                f"<td>{'YES' if t.get('smoking_gun') else ''}</td></tr>"
                for t in traj
            )
            sep = '<div style="page-break-before: always;">' if i > 0 else "<div>"
            sections.append(
                f"{sep}<h1>Trajectory: {sname}</h1>"
                f"<table><tr><th>Assignment</th><th>Submitted</th>"
                f"<th>Score</th><th>Concern</th><th>SG</th></tr>"
                f"{traj_rows}</table></div>"
            )

        return "\n".join(sections) if sections else (
            "<h1>Trajectory</h1><p>No trajectory data found.</p>"
        )

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _get_cohort_rows(self) -> List[Dict]:
        if not self._store or not self._course_id:
            return []
        rows = self._store.get_cohort(self._course_id)
        if self._student_id:
            rows = [r for r in rows
                    if str(r.get("student_id", "")) == self._student_id]
        return rows

    @staticmethod
    def _student_ids_and_names(rows: List[Dict]) -> List[tuple]:
        seen: dict = {}
        for r in rows:
            sid = str(r.get("student_id", ""))
            if sid and sid not in seen:
                seen[sid] = r.get("student_name", sid)
        return list(seen.items())

    @staticmethod
    def _safe_name(s: str) -> str:
        return "".join(
            c if c.isalnum() or c in "-_ " else "_" for c in s
        ).strip()[:40]


# ── PDF HTML template ─────────────────────────────────────────────────────────

_PDF_TEMPLATE = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{ font-family: Georgia, serif; font-size: 10pt; color: #1a1a1a; line-height: 1.5; margin: 0; }}
  h1   {{ font-size: 18pt; color: #2a1400; margin: 0 0 4pt; border-bottom: 2px solid #8B6914; padding-bottom: 4pt; }}
  h2   {{ font-size: 13pt; color: #4a2800; margin: 14pt 0 4pt; }}
  p.meta {{ font-size: 9pt; color: #555; margin: 4pt 0 12pt; }}
  p.note {{ font-size: 8pt; color: #888; font-style: italic; margin-top: 14pt; }}
  table {{ width: 100%; border-collapse: collapse; margin: 8pt 0; font-size: 9pt; }}
  th    {{ background: #2a1400; color: #F0A830; padding: 4pt 6pt; text-align: left; font-weight: bold; }}
  td    {{ padding: 3pt 6pt; border-bottom: 1px solid #d0b080; }}
  tr.total td {{ border-top: 2px solid #8B6914; font-weight: bold; background: #fdf6e3; }}
  .flag {{ color: #c00; font-weight: bold; }}
</style>
</head>
<body>{body}</body>
</html>"""
